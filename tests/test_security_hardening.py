"""The security audit's fixes (SECURITY_AUDIT.md F01-F09, F12 and the server
deploy): API keys only go over a secure connection and never to a redirect's
other host, a transcript consent belongs to the address it was given for,
keys are kept locked, spreadsheet exports can't hold formulas, uploads are
bounded, downloads are checked and marked, and Apply deletes only the
markers that were previewed. Dummy keys and fake servers only."""

import base64
import io
import json
import os
import shutil
import sys
import tempfile
import unittest
import urllib.error
import urllib.request
from unittest import mock

import _paths  # noqa: F401
from core import atomic_io, safe_http, secrets_store
from pages.manual_chat import actions, config, embedder, llm
from pages.web import download_guard
import test_manual_chat_actions as fake
import test_network_server as ns


class RedirectTests(unittest.TestCase):
    def test_a_redirect_only_stays_on_the_same_host_without_dropping_to_http(self):
        yes = safe_http.may_follow
        self.assertTrue(yes("https://api.example/v1/models", "https://api.example/v2/models"))
        self.assertTrue(yes("http://127.0.0.1:1234/a", "http://127.0.0.1:1234/b"))
        self.assertTrue(yes("http://relay.example/v1", "https://relay.example/v1"))      # up to https is fine
        self.assertFalse(yes("https://api.example/v1", "http://api.example/v1"))        # never down to http
        self.assertFalse(yes("https://api.example/v1", "https://evil.example/v1"))      # nor to another host
        self.assertFalse(yes("https://api.example/v1", "https://api.example.evil/v1"))
        self.assertFalse(yes("https://api.example/v1", "https://api.example:8443/v1"))  # nor another port
        self.assertFalse(yes("https://api.example/v1", "file:///C:/secret"))

    def test_the_handler_refuses_what_may_not_be_followed_and_keeps_the_key_off_it(self):
        handler = safe_http._SameHostOnly()
        request = urllib.request.Request("https://api.example/v1/chat",
                                         headers={"Authorization": "Bearer dummy", "x-api-key": "dummy"})
        self.assertIsNone(handler.redirect_request(request, None, 302, "Found", {}, "https://evil.example/steal"))
        self.assertIsNone(handler.redirect_request(request, None, 302, "Found", {}, "http://api.example/plain"))
        followed = handler.redirect_request(request, None, 302, "Found", {}, "https://api.example/v1/other")
        self.assertEqual(followed.full_url, "https://api.example/v1/other")


class LiveRedirectTests(unittest.TestCase):
    """Two real local servers: the first sends every request on to the second."""

    def test_a_key_does_not_reach_the_host_a_redirect_names(self):
        import http.server
        import threading
        seen = []

        class Target(http.server.BaseHTTPRequestHandler):
            def do_GET(self):
                seen.append(self.headers.get("Authorization"))
                self.send_response(200)
                self.end_headers()
                self.wfile.write(b"{}")

            def log_message(self, *args):
                pass

        target = http.server.HTTPServer(("127.0.0.1", 0), Target)

        class Mover(Target):
            def do_GET(self):
                self.send_response(302)
                # The same machine, but another host name: another origin.
                self.send_header("Location", f"http://localhost:{target.server_port}/stolen")
                self.end_headers()

            do_POST = do_GET

        mover = http.server.HTTPServer(("127.0.0.1", 0), Mover)
        for server in (target, mover):
            threading.Thread(target=server.serve_forever, daemon=True).start()
            self.addCleanup(server.server_close)
            self.addCleanup(server.shutdown)
        request = urllib.request.Request(f"http://127.0.0.1:{mover.server_port}/v1/models",
                                         headers={"Authorization": "Bearer dummy"})
        with self.assertRaises(urllib.error.HTTPError) as caught:
            safe_http.urlopen(request, timeout=5)
        self.assertEqual(caught.exception.code, 302)
        caught.exception.close()
        self.assertEqual(seen, [])                                              # the key went nowhere else
        with self.assertRaises(llm.LLMError):                                   # and chat says so in words
            llm._post(f"http://127.0.0.1:{mover.server_port}/x", {}, {"Authorization": "Bearer dummy"}, "Test", 5)
        self.assertEqual(seen, [])


class TransportTests(unittest.TestCase):
    def test_a_key_is_not_sent_to_an_http_address_off_this_pc(self):
        with mock.patch.object(llm.safe_http, "urlopen") as opened:
            with self.assertRaises(llm.LLMError):
                llm.list_openai_models("http://public.example.test/v1", "dummy")
            with self.assertRaises(llm.LLMError):
                llm.list_openai_models("http://192.168.1.20:1234/v1", "dummy")
            opened.assert_not_called()
        body = io.BytesIO(json.dumps({"data": [{"id": "m1"}]}).encode())
        with mock.patch.object(llm.safe_http, "urlopen", return_value=body):
            self.assertEqual(llm.list_openai_models("http://127.0.0.1:1234/v1", "dummy"), ["m1"])
        body = io.BytesIO(json.dumps({"data": [{"id": "m2"}]}).encode())
        with mock.patch.object(llm.safe_http, "urlopen", return_value=body):                 # no key: nothing to protect
            self.assertEqual(llm.list_openai_models("http://192.168.1.20:1234/v1"), ["m2"])

    def test_the_embedding_server_gets_the_same_rule(self):
        unsafe = embedder.ServerEmbedder("http://public.example.test/v1", "m", "dummy")
        with mock.patch.object(embedder.safe_http, "urlopen") as opened:
            ok, why = unsafe.ready()
            self.assertFalse(ok)
            self.assertIn("unencrypted", why)
            with self.assertRaises(embedder.EmbedError):
                unsafe.embed(["text"])
            opened.assert_not_called()
        self.assertEqual(embedder.ServerEmbedder("127.0.0.1:1234", "m", "dummy")._unsafe_key(), "")
        self.assertEqual(embedder.ServerEmbedder("https://embed.example/v1", "m", "dummy")._unsafe_key(), "")

    def test_a_server_is_local_by_where_it_is_not_by_what_it_is_called(self):
        for provider in (llm.PROVIDER_OLLAMA, llm.PROVIDER_LLAMACPP, llm.PROVIDER_LMSTUDIO):
            self.assertTrue(llm.LLMClient(provider, "", "m").local)                                  # its default address
            self.assertTrue(llm.LLMClient(provider, "", "m", base_url="192.168.1.20:8080").local)    # the user's network
            self.assertFalse(llm.LLMClient(provider, "", "m", base_url="https://public.example.test/v1").local)
        self.assertTrue(llm.LLMClient(llm.PROVIDER_OPENAI, "", "m", base_url="http://10.0.0.5/v1").local)
        self.assertFalse(llm.LLMClient(llm.PROVIDER_OPENAI, "", "m", base_url="https://relay.example/v1").local)

    def test_a_consent_names_the_destination(self):
        relay = llm.LLMClient(llm.PROVIDER_OPENAI, "k", "m", base_url="https://Relay.Example/v1")
        self.assertEqual(relay.destination, "relay.example")
        self.assertFalse(relay.default_address)
        gemini = llm.LLMClient(llm.PROVIDER_GEMINI, "k", "m")
        self.assertEqual((gemini.destination, gemini.default_address), ("gemini", True))
        groq = llm.LLMClient(llm.PROVIDER_GROQ, "k", "m")
        self.assertEqual((groq.destination, groq.default_address), ("api.groq.com", True))


class SecretTests(unittest.TestCase):
    @unittest.skipUnless(sys.platform == "win32", "DPAPI is Windows'")
    def test_a_key_is_sealed_and_comes_back(self):
        locked = secrets_store.lock("sk-dummy-123")
        self.assertTrue(secrets_store.is_locked(locked))
        self.assertNotIn("dummy", locked)
        self.assertEqual(secrets_store.unlock(locked), "sk-dummy-123")
        self.assertEqual(secrets_store.lock(locked), locked)                  # not locked twice
        self.assertEqual(secrets_store.unlock(secrets_store.PREFIX + "AAAA"), "")   # not this account's: asked for again

    def test_blank_and_older_values(self):
        self.assertEqual(secrets_store.lock(""), "")
        self.assertEqual(secrets_store.lock(None), "")
        self.assertEqual(secrets_store.unlock("plain-key"), "plain-key")     # saved before keys were locked
        self.assertEqual(secrets_store.unlock(None), "")

    @unittest.skipUnless(sys.platform == "win32", "DPAPI is Windows'")
    def test_stored_keys_are_locked_and_the_backup_with_them_removed(self):
        folder = tempfile.mkdtemp(prefix="buddy_keys_")
        self.addCleanup(shutil.rmtree, folder, True)

        class Settings:
            def __init__(self):
                self._path = os.path.join(folder, "settings.json")
                self.values = {"provider": "gemini", "api_key_gemini": "g-dummy", "embed_api_key": "e-dummy",
                               "api_key": "old-dummy", "api_key_groq": ""}

            def get(self, key, default=None):
                return self.values.get(key, default)

            def save(self):
                atomic_io.write_json(self._path, self.values)

        s = Settings()
        s.save()
        s.save()                                                               # the second leaves a .bak of the first
        self.assertTrue(os.path.exists(atomic_io.backup_path(s._path)))
        self.assertTrue(config.lock_stored_keys(s))
        with open(s._path, encoding="utf-8") as fh:
            on_disk = fh.read()
        for secret in ("g-dummy", "e-dummy", "old-dummy"):
            self.assertNotIn(secret, on_disk)
        self.assertFalse(os.path.exists(atomic_io.backup_path(s._path)))      # the file as it was is gone too
        self.assertEqual(config.api_key_from_settings(s, "gemini"), "g-dummy")
        self.assertEqual(secrets_store.unlock(s.get("embed_api_key")), "e-dummy")
        self.assertFalse(config.lock_stored_keys(s))                           # nothing left to do


class ExportTests(unittest.TestCase):
    def test_csv_text_that_looks_like_a_formula_is_defused(self):
        from pages.time_tracker.export_utils import safe_csv
        for text in ("=1+1", "+SUM(A1)", "-2+3", "@cmd", " =1", "\t=1", "\r=1"):
            self.assertTrue(safe_csv(text).startswith("'"), text)
        for text in ("Edit v2", "2026-10-02", "", "a=b", "Résumé"):
            self.assertEqual(safe_csv(text), text)
        self.assertEqual(safe_csv(5), 5)

    def test_xlsx_stores_text_as_text(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("openpyxl isn't installed")
        from pages.time_tracker.export_utils import append_text_safe
        ws = Workbook().active
        append_text_safe(ws, ["=1+1", "plain", 3])
        self.assertEqual([c.data_type for c in ws[1]], ["s", "s", "n"])
        self.assertEqual(ws["A1"].value, "=1+1")


class UploadTests(ns.Harness):
    def part(self, session, image_id, seq, data, **extra):
        return self.request(session, type="image_part", id=image_id, seq=seq, data=data, **extra)

    def test_an_empty_piece_is_refused(self):
        a = self.user("Ann")
        image_id = "a" * 32
        self.assertIsNone(self.part(a, image_id, 0, ns.b64(30)))
        self.assertEqual(self.part(a, image_id, 1, "")["code"], "bad_image")
        self.assertIsNone(a.upload)

    def test_there_is_a_limit_to_how_many_pieces(self):
        a = self.user("Ann")
        image_id = "b" * 32
        for seq in range(core_parts()):
            self.assertIsNone(self.part(a, image_id, seq, ns.b64(4)), seq)
        self.assertEqual(self.part(a, image_id, core_parts(), ns.b64(4))["code"], "bad_image")
        self.assertIsNone(a.upload)

    def test_an_upload_that_takes_too_long_is_dropped(self):
        a = self.user("Ann")
        image_id = "c" * 32
        self.assertIsNone(self.part(a, image_id, 0, ns.b64(30)))
        self.clock.now += ns.core.UPLOAD_SECONDS + 1
        self.assertEqual(self.part(a, image_id, 1, ns.b64(30))["code"], "bad_image")

    def test_screenshots_have_the_same_limits(self):
        a = ns.FakeSession()
        image_id = "d" * 32
        self.request(a, type="bug_part", id=image_id, seq=0, data=ns.b64(30))
        self.assertEqual(self.request(a, type="bug_part", id=image_id, seq=1, data="")["code"], "bad_image")
        b = ns.FakeSession()
        for seq in range(core_parts()):
            self.request(b, type="bug_part", id=image_id, seq=seq, data=ns.b64(4))
        self.assertEqual(self.request(b, type="bug_part", id=image_id, seq=core_parts(), data=ns.b64(4))["code"],
                         "bad_image")


def core_parts():
    return ns.core.MAX_IMAGE_PARTS


class ImageImporterTests(unittest.TestCase):
    PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")

    def test_a_download_is_saved_as_what_it_is(self):
        try:
            from pages.image_importer import clipboard
        except ImportError:
            self.skipTest("PySide6 isn't installed")
        self.assertEqual(clipboard.verified_extension(self.PNG), ".png")
        for bad in (b"MZ\x90\x00 a program", b"", b"<svg xmlns='http://www.w3.org/2000/svg'/>", b"GIF" + b"\x00" * 3):
            with self.assertRaises(ValueError):
                clipboard.verified_extension(bad)


class MarkerTests(unittest.TestCase):
    class Timeline(fake.FakeTimeline):
        def DeleteMarkerAtFrame(self, frame):
            return self.markers.pop(frame, None) is not None

        def DeleteMarkersByColor(self, color):
            raise AssertionError("a colour must never be what is deleted by")

    def setup(self, markers):
        timeline = self.Timeline("Edit", "tl-1")
        timeline.markers = dict(markers)
        project = fake.FakeProject("Show", "pr-1", timeline)
        return fake.FakeController(project), timeline

    def red(self, n):
        return {frame: {"color": "Red", "name": f"m{frame}", "note": ""} for frame in range(n)}

    def test_apply_deletes_only_the_markers_that_were_shown(self):
        controller, timeline = self.setup({10: {"color": "Red", "name": "a", "note": ""}})
        proposal = actions.preview(controller, "delete_markers", {"color": "Red"})
        timeline.markers[20] = {"color": "Red", "name": "added after", "note": ""}      # not in the preview
        actions.execute(controller, proposal)
        self.assertEqual(list(timeline.markers), [20])

    def test_a_marker_that_changed_since_is_left_alone(self):
        controller, timeline = self.setup({10: {"color": "Red", "name": "a", "note": ""},
                                           11: {"color": "Red", "name": "b", "note": ""}})
        proposal = actions.preview(controller, "delete_markers", {"color": "Red"})
        timeline.markers[10]["name"] = "renamed"
        timeline.markers[11]["color"] = "Blue"
        result = actions.execute(controller, proposal)
        self.assertEqual(sorted(timeline.markers), [10, 11])
        self.assertIn("Deleted 0 of 2", result)

    def test_the_item_limit_covers_a_colour_too(self):
        controller, _ = self.setup(self.red(actions.MAX_ITEMS_PER_ACTION + 1))
        with self.assertRaises(actions.ActionError):
            actions.preview(controller, "delete_markers", {"color": "Red"})
        controller, _ = self.setup(self.red(actions.MAX_ITEMS_PER_ACTION))
        self.assertEqual(len(actions.preview(controller, "delete_markers", {"color": "Red"}).details),
                         actions.MAX_ITEMS_PER_ACTION)


class DownloadGuardTests(unittest.TestCase):
    def test_programs_and_scripts(self):
        for name in ("setup.exe", "Setup.EXE", "run.bat", "x.ps1", "report.pdf.exe", "a.msi", "link.lnk", "p.js"):
            self.assertTrue(download_guard.is_program_file(name), name)
        for name in ("clip.mp4", "notes.txt", "setup.exe.txt", "exe", "archive.zip", "photo.png", ""):
            self.assertFalse(download_guard.is_program_file(name), name)

    def test_room_on_the_drive(self):
        gib = 1024 ** 3
        self.assertTrue(download_guard.room_for(10 * gib, 2 * gib))
        self.assertFalse(download_guard.room_for(10 * gib, 9.5 * gib))
        self.assertFalse(download_guard.room_for(gib // 2, 0))
        self.assertTrue(download_guard.room_for(10 * gib, 0))                  # size not known yet

    def test_a_burst_from_one_site_needs_the_users_ok_once(self):
        guard = download_guard.BurstGuard()
        results = [guard.burst("a.example", 100.0 + i) for i in range(download_guard.BURST_COUNT + 1)]
        self.assertEqual(results, [False] * download_guard.BURST_COUNT + [True])
        self.assertFalse(guard.burst("b.example", 105.0))                        # another site
        guard.allow("a.example", 106.0)
        self.assertFalse(guard.burst("a.example", 107.0))
        self.assertTrue(any(guard.burst("a.example", 500.0 + i) for i in range(download_guard.BURST_COUNT + 1)))

    def test_a_slow_trickle_is_not_a_burst(self):
        guard = download_guard.BurstGuard()
        self.assertFalse(any(guard.burst("a.example", i * 30.0) for i in range(20)))

    @unittest.skipUnless(sys.platform == "win32", "Zone.Identifier is Windows'")
    def test_a_finished_file_is_marked_as_from_the_internet(self):
        folder = tempfile.mkdtemp(prefix="buddy_zone_")
        self.addCleanup(shutil.rmtree, folder, True)
        path = os.path.join(folder, "file.bin")
        open(path, "wb").write(b"x")
        if not download_guard.mark_from_internet(path, "https://cdn.example/file.bin", "https://site.example/page"):
            self.skipTest("this drive has no alternate data streams")
        text = open(path + ":Zone.Identifier", encoding="ascii").read()
        self.assertIn("ZoneId=3", text)
        self.assertIn("HostUrl=https://cdn.example/file.bin", text)
        self.assertTrue(download_guard.mark_from_internet(path))               # already marked: left as it is
        self.assertFalse(download_guard.mark_from_internet(os.path.join(folder, "missing.bin")))

    def test_a_marker_url_cannot_add_lines(self):
        if sys.platform != "win32":
            self.skipTest("Zone.Identifier is Windows'")
        folder = tempfile.mkdtemp(prefix="buddy_zone_")
        self.addCleanup(shutil.rmtree, folder, True)
        path = os.path.join(folder, "f.bin")
        open(path, "wb").write(b"x")
        if not download_guard.mark_from_internet(path, "https://x.example/a\r\nZoneId=1"):
            self.skipTest("this drive has no alternate data streams")
        with open(path + ":Zone.Identifier", encoding="ascii") as fh:
            lines = fh.read().splitlines()
        self.assertEqual([l for l in lines if l.startswith("ZoneId")], ["ZoneId=3"])


if __name__ == "__main__":
    unittest.main()
