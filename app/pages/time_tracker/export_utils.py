#!/usr/bin/env python3
"""
Export helpers for the "This Project"/"History" tabs and Reports - CSV,
JSON, PDF, and XLSX, all taking the same (entries, data_mgr, path) shape so
the page can dispatch to whichever one the user picked with a single
call.
"""

import csv
import json

from .ui_utils import format_date, format_clock, format_duration_words

COLUMN_HEADERS = ["Date", "Project", "Start", "End", "Duration", "Notes"]

# A spreadsheet reads a cell starting with one of these as a formula - and
# project names and notes are the user's (or an imported backup's) own text.
_FORMULA_START = ("=", "+", "-", "@")


def safe_csv(value):
    """Text for a CSV cell that a spreadsheet shows as written: one that
    starts like a formula (after any spaces and control characters, which
    some spreadsheets skip) gets a leading apostrophe."""
    if not isinstance(value, str):
        return value
    first = value.lstrip(" \t\r\n\x00\x0b\x0c")
    if value[:1] in ("\t", "\r", "\n") or first.startswith(_FORMULA_START):
        return "'" + value
    return value


def append_text_safe(ws, row):
    """Adds `row` to a worksheet with every text cell stored as text, never
    as a formula (openpyxl turns a string starting "=" into one)."""
    ws.append(row)
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def _row_values(entry, data_mgr):
    return [
        format_date(entry["start"]),
        entry.get("project", ""),
        format_clock(entry["start"]),
        format_clock(entry.get("end")),
        format_duration_words(data_mgr.duration_seconds(entry)),
        entry.get("notes", ""),
    ]


def export_csv(entries, data_mgr, path):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["Date", "Project", "Start", "End", "Duration (seconds)", "Notes"])
        for entry in entries:
            writer.writerow([
                format_date(entry["start"]),
                safe_csv(entry.get("project", "")),
                entry["start"],
                entry.get("end", ""),
                int(data_mgr.duration_seconds(entry)),
                safe_csv(entry.get("notes", "")),
            ])


def export_json(entries, data_mgr, path):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(entries, f, indent=2)


def export_xlsx(entries, data_mgr, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Time entries"

    ws.append(COLUMN_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for entry in entries:
        append_text_safe(ws, _row_values(entry, data_mgr))

    widths = [12, 32, 10, 10, 12, 40]
    for col_index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_index).column_letter].width = width

    wb.save(path)


def export_pdf(entries, data_mgr, path):
    # Built on Qt's own QtPrintSupport/QTextDocument - no extra dependency
    # needed (unlike XLSX, which needs openpyxl), since Qt can already
    # render HTML to a PDF file directly.
    from PySide6.QtCore import QMarginsF
    from PySide6.QtGui import QPageLayout, QPageSize, QTextDocument
    from PySide6.QtPrintSupport import QPrinter

    import html as pyhtml
    rows_html = "".join(
        "<tr>" + "".join(f"<td>{pyhtml.escape(str(value))}</td>" for value in _row_values(entry, data_mgr)) + "</tr>"
        for entry in entries
    )
    header_html = "".join(f"<th>{h}</th>" for h in COLUMN_HEADERS)
    html = f"""
    <html><head><style>
        body {{ font-family: Segoe UI, Arial, sans-serif; font-size: 10pt; }}
        h1 {{ font-size: 16pt; }}
        table {{ border-collapse: collapse; width: 100%; }}
        th, td {{ border: 1px solid #999; padding: 4px 8px; text-align: left; }}
        th {{ background-color: #eee; }}
    </style></head><body>
    <h1>Time Tracker – Time entries</h1>
    <table><tr>{header_html}</tr>{rows_html}</table>
    </body></html>
    """

    document = QTextDocument()
    document.setHtml(html)

    printer = QPrinter(QPrinter.HighResolution)
    printer.setOutputFormat(QPrinter.PdfFormat)
    printer.setOutputFileName(path)
    printer.setPageSize(QPageSize(QPageSize.A4))
    printer.setPageMargins(QMarginsF(15, 15, 15, 15), QPageLayout.Millimeter)
    document.print_(printer)


EXPORTERS = {
    "CSV": {"func": export_csv, "filename": "time_tracker_export.csv", "filter": "CSV files (*.csv)"},
    "JSON": {"func": export_json, "filename": "time_tracker_export.json", "filter": "JSON files (*.json)"},
    "PDF": {"func": export_pdf, "filename": "time_tracker_export.pdf", "filter": "PDF files (*.pdf)"},
    "Spreadsheet (XLSX)": {
        "func": export_xlsx, "filename": "time_tracker_export.xlsx",
        "filter": "Excel files (*.xlsx)",
    },
}


# --- Reports tab exports ---
# report_data shape (see reports.py's build()):
#   {"today_label": str, "today_seconds": float,
#    "today_earnings_by_currency": {"USD": float, ...},
#    "week_label": str, "week_seconds": float,
#    "week_earnings_by_currency": {"USD": float, ...},
#    "month_label": str, "month_seconds": float,
#    "month_earnings_by_currency": {"USD": float, ...},
#    "daily": [(label, seconds), ...],
#    "by_project": [(project, seconds, earnings, currency), ...],
#    optionally "custom_label"/"custom_seconds"/"custom_earnings_by_currency"
#    - present only once the user has actually run a Custom Range
#    calculation this session (see reports.custom_range)}
# earnings is 0.0 for a project with no rate set - always present so exports
# don't need to special-case it. Earnings are kept split by currency
# throughout, never summed across currencies - see currencies.py.
_BASE_TOTALS_SECTIONS = [
    ("today_label", "today_seconds", "today_earnings_by_currency"),
    ("week_label", "week_seconds", "week_earnings_by_currency"),
    ("month_label", "month_seconds", "month_earnings_by_currency"),
]


def _totals_sections_for(report_data):
    sections = list(_BASE_TOTALS_SECTIONS)
    if "custom_label" in report_data:
        sections.append(("custom_label", "custom_seconds", "custom_earnings_by_currency"))
    return sections


def _format_earnings(report_data, currency_key, formatter):
    by_currency = report_data.get(currency_key, {})
    return [formatter(amount, currency) for currency, amount in sorted(by_currency.items())]


def export_report_csv(report_data, path):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["Metric", "Duration (seconds)"])
        for label_key, seconds_key, currency_key in _totals_sections_for(report_data):
            writer.writerow([safe_csv(report_data[label_key]), int(report_data[seconds_key])])
            for currency, amount in sorted(report_data.get(currency_key, {}).items()):
                writer.writerow([safe_csv(f"{report_data[label_key]} Earnings ({currency})"), f"{amount:.2f}"])
        writer.writerow([])
        writer.writerow(["Date", "Duration (seconds)"])
        for label, seconds in report_data["daily"]:
            writer.writerow([safe_csv(label), int(seconds)])
        writer.writerow([])
        writer.writerow(["Project", "Duration (seconds)", "Earnings", "Currency"])
        for project, seconds, earnings, currency in report_data["by_project"]:
            writer.writerow([safe_csv(project), int(seconds), f"{earnings:.2f}", safe_csv(currency)])


def export_report_json(report_data, path):
    payload = {}
    for label_key, seconds_key, currency_key in _totals_sections_for(report_data):
        payload[label_key] = report_data[label_key]
        payload[seconds_key] = int(report_data[seconds_key])
        payload[currency_key.replace("_by_currency", "")] = {
            currency: round(amount, 2) for currency, amount in report_data.get(currency_key, {}).items()
        }
    payload["daily"] = [{"label": label, "seconds": int(s)} for label, s in report_data["daily"]]
    payload["by_project"] = [
        {"project": p, "seconds": int(s), "earnings": round(e, 2), "currency": c}
        for p, s, e, c in report_data["by_project"]
    ]
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2)


def export_report_xlsx(report_data, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from .currencies import format_money

    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    summary.append(["Metric", "Duration"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    for label_key, seconds_key, currency_key in _totals_sections_for(report_data):
        append_text_safe(summary, [report_data[label_key], format_duration_words(report_data[seconds_key])])
        for text in _format_earnings(report_data, currency_key, lambda a, c: f"{c} Earnings: {format_money(a, c)}"):
            append_text_safe(summary, [text])
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 16

    daily = wb.create_sheet("Last 7 days")
    daily.append(["Date", "Duration"])
    daily["A1"].font = daily["B1"].font = Font(bold=True)
    for label, seconds in report_data["daily"]:
        append_text_safe(daily, [label, format_duration_words(seconds) if seconds else "-"])
    daily.column_dimensions["A"].width = 16
    daily.column_dimensions["B"].width = 16

    projects = wb.create_sheet("By project")
    projects.append(["Project", "Duration", "Earnings"])
    for cell in projects[1]:
        cell.font = Font(bold=True)
    for project, seconds, earnings, currency in report_data["by_project"]:
        earnings_text = format_money(earnings, currency) if earnings > 0 else ""
        append_text_safe(projects, [project, format_duration_words(seconds), earnings_text])
    projects.column_dimensions["A"].width = 32
    projects.column_dimensions["B"].width = 16
    projects.column_dimensions["C"].width = 16

    wb.save(path)


def export_report_pdf(report_data, path):
    from PySide6.QtCore import QMarginsF
    from PySide6.QtGui import QPageLayout, QPageSize, QTextDocument
    from PySide6.QtPrintSupport import QPrinter
    from .currencies import format_money

    import html as pyhtml
    def table_html(headers, rows):
        header_html = "".join(f"<th>{pyhtml.escape(str(h))}</th>" for h in headers)
        rows_html = "".join(
            "<tr>" + "".join(f"<td>{pyhtml.escape(str(v))}</td>" for v in row) + "</tr>" for row in rows
        )
        return f"<table><tr>{header_html}</tr>{rows_html}</table>"

    daily_rows = [(label, format_duration_words(s) if s else "-") for label, s in report_data["daily"]]
    project_rows = [
        (project, format_duration_words(s), format_money(e, c) if e > 0 else "-")
        for project, s, e, c in report_data["by_project"]
    ]

    totals_html = ""
    for label_key, seconds_key, currency_key in _totals_sections_for(report_data):
        earnings_lines = _format_earnings(report_data, currency_key, lambda a, c: f"{format_money(a, c)} earned")
        earnings_html = "".join(f'<p style="font-size: 12pt;">{line}</p>' for line in earnings_lines)
        totals_html += f"""
        <h2>{report_data[label_key]}</h2>
        <p style="font-size: 18pt; font-weight: bold;">{format_duration_words(report_data[seconds_key])}</p>
        {earnings_html}
        """

    html = f"""
    <html><head><style>
        body {{ font-family: Segoe UI, Arial, sans-serif; font-size: 10pt; }}
        h1 {{ font-size: 16pt; }}
        h2 {{ font-size: 13pt; margin-top: 24px; }}
        table {{ border-collapse: collapse; width: 100%; }}
        th, td {{ border: 1px solid #999; padding: 4px 8px; text-align: left; }}
        th {{ background-color: #eee; }}
    </style></head><body>
    <h1>Time Tracker – Reports</h1>
    {totals_html}
    <h2>Last 7 days</h2>
    {table_html(["Date", "Duration"], daily_rows)}
    <h2>By project (all time)</h2>
    {table_html(["Project", "Duration", "Earnings"], project_rows)}
    </body></html>
    """

    document = QTextDocument()
    document.setHtml(html)

    printer = QPrinter(QPrinter.HighResolution)
    printer.setOutputFormat(QPrinter.PdfFormat)
    printer.setOutputFileName(path)
    printer.setPageSize(QPageSize(QPageSize.A4))
    printer.setPageMargins(QMarginsF(15, 15, 15, 15), QPageLayout.Millimeter)
    document.print_(printer)


# --- Invoice export (Reports tab > Generate Invoice) ---
# Scoped to one project + date range, using that project's own rate/currency
# - unlike the general History/Reports exports, this is meant to be handed
# to a client, so line items are grouped per day worked rather than per raw
# tracking session.
def export_invoice_pdf(entries, data_mgr, project_name, start_date, end_date, path,
                        bill_to="", invoice_number=""):
    from collections import defaultdict
    import html as pyhtml
    from PySide6.QtCore import QMarginsF
    from PySide6.QtGui import QPageLayout, QPageSize, QTextDocument
    from PySide6.QtPrintSupport import QPrinter
    from .currencies import format_money

    info = data_mgr.get_project_rate_info(project_name)
    rate, currency = info["rate"], info["currency"]

    by_day_seconds = defaultdict(float)
    by_day_notes = defaultdict(list)
    for entry in entries:
        if entry.get("project") != project_name or entry.get("hidden"):
            continue
        entry_date = entry["start"][:10]
        if not (start_date <= entry_date <= end_date):
            continue
        by_day_seconds[entry_date] += data_mgr.duration_seconds(entry)
        note = entry.get("notes", "").strip()
        if note and note not in by_day_notes[entry_date]:
            by_day_notes[entry_date].append(note)

    days = sorted(by_day_seconds.keys())
    total_seconds = sum(by_day_seconds.values())
    total_hours = total_seconds / 3600.0
    total_amount = total_hours * rate

    def row_html(day):
        hours = by_day_seconds[day] / 3600.0
        description = "; ".join(by_day_notes[day]) or "—"
        if len(description) > 80:
            description = description[:79] + "…"
        amount = format_money(hours * rate, currency, short=True) if rate > 0 else "—"
        return f"<tr><td>{pyhtml.escape(day)}</td><td>{pyhtml.escape(description)}</td><td>{hours:.2f}</td><td>{amount}</td></tr>"

    rows_html = "".join(row_html(day) for day in days)
    bill_to_html = f'<p><strong>Bill To:</strong> {pyhtml.escape(bill_to)}</p>' if bill_to.strip() else ""
    invoice_number_html = f'<p><strong>Invoice #:</strong> {pyhtml.escape(invoice_number)}</p>' if invoice_number.strip() else ""
    rate_html = (
        f"<p>Rate: {format_money(rate, currency)} / hour</p>" if rate > 0
        else "<p>No rate set for this project – amounts left blank.</p>"
    )

    html = f"""
    <html><head><style>
        body {{ font-family: Segoe UI, Arial, sans-serif; font-size: 10pt; }}
        h1 {{ font-size: 18pt; margin-bottom: 2px; }}
        .sub {{ color: #555; margin-top: 0; }}
        table {{ border-collapse: collapse; width: 100%; margin-top: 16px; }}
        th, td {{ border: 1px solid #999; padding: 4px 8px; text-align: left; }}
        th {{ background-color: #eee; }}
        td:nth-child(3), th:nth-child(3), td:nth-child(4), th:nth-child(4) {{ text-align: right; }}
        .totals {{ margin-top: 16px; font-size: 12pt; }}
        .totals strong {{ font-size: 15pt; }}
    </style></head><body>
    <h1>Invoice</h1>
    <p class="sub">{pyhtml.escape(project_name)} — {start_date} to {end_date}</p>
    {invoice_number_html}
    {bill_to_html}
    {rate_html}
    <table>
        <tr><th>Date</th><th>Description</th><th>Hours</th><th>Amount</th></tr>
        {rows_html if rows_html else '<tr><td colspan="4">No tracked time in this date range.</td></tr>'}
    </table>
    <div class="totals">
        <p>Total time: {format_duration_words(total_seconds)} ({total_hours:.2f} hours)</p>
        <p><strong>{"Total due: " + format_money(total_amount, currency) if rate > 0 else ""}</strong></p>
    </div>
    </body></html>
    """

    document = QTextDocument()
    document.setHtml(html)

    printer = QPrinter(QPrinter.HighResolution)
    printer.setOutputFormat(QPrinter.PdfFormat)
    printer.setOutputFileName(path)
    printer.setPageSize(QPageSize(QPageSize.A4))
    printer.setPageMargins(QMarginsF(15, 15, 15, 15), QPageLayout.Millimeter)
    document.print_(printer)


REPORT_EXPORTERS = {
    "PDF": {"func": export_report_pdf, "filename": "time_tracker_report.pdf", "filter": "PDF files (*.pdf)"},
    "CSV": {"func": export_report_csv, "filename": "time_tracker_report.csv", "filter": "CSV files (*.csv)"},
    "JSON": {"func": export_report_json, "filename": "time_tracker_report.json", "filter": "JSON files (*.json)"},
    "Spreadsheet (XLSX)": {
        "func": export_report_xlsx, "filename": "time_tracker_report.xlsx",
        "filter": "Excel files (*.xlsx)",
    },
}
